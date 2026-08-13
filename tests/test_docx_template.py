from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement

from vehicle_valuation.docx_template import (
    fill_docx_template,
)


def test_multiline_placeholder_creates_real_paragraphs(
    tmp_path: Path,
) -> None:
    """RAG多段内容应写成多个Word段落。"""

    template_path = tmp_path / "template.docx"
    document = Document()
    document.add_paragraph("{{VALUATION_METHOD}}")
    document.save(template_path)

    output = fill_docx_template(
        template_path,
        {
            "{{VALUATION_METHOD}}": (
                "第一段。\n\n第二段。\n实例1：测试案例。"
            )
        },
    )
    output_path = tmp_path / "output.docx"
    output_path.write_bytes(output)

    result = Document(output_path)

    assert [
        paragraph.text
        for paragraph in result.paragraphs
    ] == [
        "第一段。",
        "第二段。",
        "实例1：测试案例。",
    ]


def test_stale_toc_page_numbers_are_removed(
    tmp_path: Path,
) -> None:
    """模板目录不应保留与新报告不一致的旧页码。"""

    template_path = tmp_path / "template.docx"
    document = Document()
    document.add_paragraph("目   录")
    document.add_paragraph("一、评估目的\t........\t9")
    document.add_paragraph("声   明")
    document.save(template_path)

    output = fill_docx_template(template_path, {})
    output_path = tmp_path / "output.docx"
    output_path.write_bytes(output)

    result = Document(output_path)

    assert [paragraph.text for paragraph in result.paragraphs] == [
        "目   录",
        "一、评估目的",
        "声   明",
    ]


def test_empty_numbered_template_paragraphs_are_removed(
    tmp_path: Path,
) -> None:
    """清空正文后遗留的编号不应出现在最终报告。"""

    template_path = tmp_path / "template.docx"
    document = Document()
    document.add_paragraph("六、评估依据")
    document.add_paragraph("{{VALUATION_BASIS}}")
    empty_numbered = document.add_paragraph("")
    properties = empty_numbered._p.get_or_add_pPr()
    numbering = OxmlElement("w:numPr")
    numbering.append(OxmlElement("w:ilvl"))
    numbering.append(OxmlElement("w:numId"))
    properties.append(numbering)
    document.add_paragraph("")
    document.add_paragraph("七、评估方法")
    document.save(template_path)

    output = fill_docx_template(
        template_path,
        {"{{VALUATION_BASIS}}": "准则依据。"},
    )
    output_path = tmp_path / "output.docx"
    output_path.write_bytes(output)

    result = Document(output_path)

    assert [paragraph.text for paragraph in result.paragraphs] == [
        "六、评估依据",
        "准则依据。",
        "",
        "七、评估方法",
    ]
