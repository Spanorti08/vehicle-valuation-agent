from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from vehicle_valuation.loaders import load_synthetic_valuation_request
from vehicle_valuation.model import MarketListing
from vehicle_valuation.output_validation import validate_output_consistency


DATA_DIR = Path(__file__).parents[1] / "data" / "synthetic"


def test_output_consistency_accepts_matching_artifacts() -> None:
    request = load_synthetic_valuation_request(DATA_DIR, "2026-06-30")
    listing = MarketListing(
        source_url="https://example.com/case-1",
        vehicle_model="测试车型",
        registration_year=2018,
        mileage_km=80_000,
        city="北京",
        price_cny=Decimal("80000"),
    )
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title in ("封面", "汇总表", "固定资产汇总", "车辆信息", "计算表", "案例1"):
        workbook.create_sheet(title)
    workbook["车辆信息"]["B7"] = request.subject_vehicle.asset_id
    workbook["车辆信息"]["C7"] = request.subject_vehicle.plate_number
    workbook["车辆信息"]["D7"] = request.subject_vehicle.vehicle_name
    workbook["计算表"]["AA2"] = 80000
    workbook["案例1"]["B1"] = str(listing.source_url)
    workbook["案例1"]["I13"] = 80000
    excel_buffer = BytesIO()
    workbook.save(excel_buffer)

    document = Document()
    document.add_paragraph(
        f"车辆{request.subject_vehicle.plate_number}评估值80,000元"
    )
    report_buffer = BytesIO()
    document.save(report_buffer)

    report = validate_output_consistency(
        excel_buffer.getvalue(),
        report_buffer.getvalue(),
        request,
        [listing],
        Decimal("80000"),
        {str(listing.source_url): (b"top", b"detail")},
        {"chunk-1"},
        {"chunk-1"},
    )

    assert report.passed is True
    assert report.issues == []


def test_output_consistency_rejects_unknown_citation_and_placeholder() -> None:
    request = load_synthetic_valuation_request(DATA_DIR, "2026-06-30")
    workbook = Workbook()
    workbook["Sheet"].title = "计算表"
    excel_buffer = BytesIO()
    workbook.save(excel_buffer)
    document = Document()
    document.add_paragraph("{{CLIENT_NAME}}")
    report_buffer = BytesIO()
    document.save(report_buffer)

    report = validate_output_consistency(
        excel_buffer.getvalue(),
        report_buffer.getvalue(),
        request,
        [],
        Decimal("80000"),
        {},
        {"missing"},
        {"valid"},
    )

    assert report.passed is False
    assert any(issue.check == "RAG引用" for issue in report.issues)
    assert any(issue.check == "模板占位符" for issue in report.issues)
