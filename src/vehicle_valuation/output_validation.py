"""Excel、Word、案例证据与 RAG 引用的最终一致性检查。"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from docx import Document
from openpyxl import load_workbook
from pydantic import BaseModel, Field

from vehicle_valuation.model import MarketListing, ValuationRequest


class OutputConsistencyIssue(BaseModel):
    check: str
    message: str
    auto_fixable: bool = False


class OutputConsistencyReport(BaseModel):
    passed: bool
    issues: list[OutputConsistencyIssue] = Field(default_factory=list)


def _document_text(document: Document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def validate_output_consistency(
    excel_bytes: bytes,
    report_bytes: bytes,
    request: ValuationRequest,
    listings: list[MarketListing],
    final_value: Decimal,
    screenshots: dict[str, tuple[bytes, bytes]],
    cited_chunk_ids: set[str],
    valid_chunk_ids: set[str],
    forbidden_terms: tuple[str, ...] = (),
) -> OutputConsistencyReport:
    """执行只读硬校验；非白名单内容不进行自动改写。"""

    issues: list[OutputConsistencyIssue] = []
    workbook = load_workbook(BytesIO(excel_bytes), data_only=False)
    required = {
        "封面",
        "汇总表",
        "固定资产汇总",
        "车辆信息",
        "计算表",
        *{f"案例{index}" for index in range(1, len(listings) + 1)},
    }
    missing = required - set(workbook.sheetnames)
    if missing:
        issues.append(OutputConsistencyIssue(
            check="Excel页签",
            message="缺少页签：" + "、".join(sorted(missing)),
        ))
    if "车辆信息" in workbook.sheetnames:
        vehicle_sheet = workbook["车辆信息"]
        expected = request.subject_vehicle
        for cell, value, label in (
            ("B7", expected.asset_id, "资产编号"),
            ("C7", expected.plate_number, "车牌号"),
            ("D7", expected.vehicle_name, "车辆名称"),
        ):
            if str(vehicle_sheet[cell].value) != str(value):
                issues.append(OutputConsistencyIssue(
                    check="Excel车辆信息",
                    message=f"{label}与结构化输入不一致",
                ))
    if "计算表" in workbook.sheetnames:
        recorded = workbook["计算表"]["AA2"].value
        if recorded is None or Decimal(str(recorded)) != final_value:
            issues.append(OutputConsistencyIssue(
                check="估值金额",
                message="Excel中的 Python 复核值与计算结果不一致",
            ))
    for index, listing in enumerate(listings, start=1):
        sheet_name = f"案例{index}"
        url = str(listing.source_url)
        if sheet_name in workbook.sheetnames:
            case_sheet = workbook[sheet_name]
            if str(case_sheet["B1"].value) != url:
                issues.append(OutputConsistencyIssue(
                    check="案例链接",
                    message=f"{sheet_name}链接与确认案例不一致",
                ))
            if Decimal(str(case_sheet["I13"].value)) != listing.price_cny:
                issues.append(OutputConsistencyIssue(
                    check="案例价格",
                    message=f"{sheet_name}价格与同次网页读取结果不一致",
                ))
        images = screenshots.get(url)
        if images is None or not all(images):
            issues.append(OutputConsistencyIssue(
                check="案例截图",
                message=f"案例{index}缺少同次读取截图",
            ))

    report_text = _document_text(Document(BytesIO(report_bytes)))
    normalized_report_text = report_text.replace(",", "")
    for value, label in (
        (request.subject_vehicle.plate_number, "车牌号"),
        (f"{final_value:.0f}", "评估值"),
    ):
        if str(value).replace(",", "") not in normalized_report_text:
            issues.append(OutputConsistencyIssue(
                check="Word事实",
                message=f"Word未找到已确认的{label}",
            ))
    if "{{" in report_text or "}}" in report_text:
        issues.append(OutputConsistencyIssue(
            check="模板占位符",
            message="Word仍存在未替换占位符",
            auto_fixable=True,
        ))
    for term in forbidden_terms:
        if term and term in report_text:
            issues.append(OutputConsistencyIssue(
                check="模板污染",
                message=f"Word仍包含旧项目内容：{term}",
            ))
    unknown_citations = cited_chunk_ids - valid_chunk_ids
    if unknown_citations:
        issues.append(OutputConsistencyIssue(
            check="RAG引用",
            message="存在无效证据ID：" + "、".join(sorted(unknown_citations)),
        ))
    return OutputConsistencyReport(passed=not issues, issues=issues)
