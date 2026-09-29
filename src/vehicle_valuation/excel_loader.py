from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from vehicle_valuation.model import SubjectVehicle


def normalize_year_month(value: object) -> str:
    """把Excel中的日期统一转换为YYYY-MM格式。"""

    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m")

    text = str(value).strip()

    for date_format in ("%Y-%m", "%Y/%m", "%Y年%m月"):
        try:
            parsed_date = datetime.strptime(text, date_format)
            return parsed_date.strftime("%Y-%m")
        except ValueError:
            continue

    raise ValueError(f"无法识别年月格式：{value}")


def load_subject_vehicles_from_excel(
    excel_path: Path,
) -> list[SubjectVehicle]:
    """从车辆sheet读取全部待评估车辆。"""

    workbook = load_workbook(
        excel_path,
        data_only=True,
    )
    worksheet = workbook["车辆"]

    vehicles = []

    for row_number in range(7, worksheet.max_row + 1):
        asset_id = worksheet.cell(row_number, 2).value

        if asset_id is None:
            continue

        vehicle = SubjectVehicle(
            sequence_number=worksheet.cell(row_number, 1).value,
            asset_id=str(asset_id),
            plate_number=str(
                worksheet.cell(row_number, 3).value
            ).strip(),
            vehicle_name=str(
                worksheet.cell(row_number, 4).value
            ).strip(),
            manufacturer=str(
                worksheet.cell(row_number, 6).value
            ).strip(),
            vin=(
                str(worksheet.cell(row_number, 25).value).strip()
                if worksheet.cell(row_number, 25).value
                else None
            ),
            unit=worksheet.cell(row_number, 7).value,
            quantity=worksheet.cell(row_number, 8).value,
            purchase_date=normalize_year_month(
                worksheet.cell(row_number, 9).value
            ),
            in_service_date=normalize_year_month(
                worksheet.cell(row_number, 10).value
            ),
            mileage_km=worksheet.cell(row_number, 11).value,
            book_value_original_cny=worksheet.cell(
                row_number, 16
            ).value,
            book_value_net_cny=worksheet.cell(
                row_number, 17
            ).value,
        )

        vehicles.append(vehicle)

    return vehicles
