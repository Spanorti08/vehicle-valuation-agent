import json
import re
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader


def extract_pdf_text(path: Path) -> str:
    """读取PDF每一页的文字，并合并成完整文本。"""
    reader = PdfReader(path)

    pages = [
        page.extract_text() or ""
        for page in reader.pages
    ]

    return "\n".join(pages).strip()


def extract_docx_text(path: Path) -> str:
    """读取Word中的段落和表格文字。"""
    document = Document(path)

    lines = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if cells:
                lines.append(" | ".join(cells))

    return "\n".join(lines).strip()


def extract_document_text(path: Path) -> str:
    """根据文件扩展名选择PDF或Word解析方式。"""
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return extract_pdf_text(path)

    if suffix == ".docx":
        return extract_docx_text(path)

    raise ValueError(f"暂不支持该文件类型：{suffix}")


def load_source_manifest(path: Path) -> list[dict[str, Any]]:
    """读取sources.json中的资料目录。"""
    data = json.loads(path.read_text(encoding="utf-8"))

    if not isinstance(data, list):
        raise ValueError("sources.json最外层必须是列表")

    return data


def build_extracted_documents(
    raw_dir: Path,
    sources_path: Path,
) -> list[dict[str, Any]]:
    """按照资料目录解析所有原始文件。"""
    sources = load_source_manifest(sources_path)
    documents = []

    for source in sources:
        document_path = raw_dir / source["filename"]

        if not document_path.exists():
            raise FileNotFoundError(
                f"找不到资料：{document_path}"
            )

        content = extract_document_text(document_path)

        documents.append(
            {
                **source,
                "content": content,
                "character_count": len(content),
            }
        )

    return documents


def build_article_chunks(
    documents: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """按照准则中的条款标记，将文档切分成检索片段。"""
    chunks = []

    article_pattern = (
        r"(?=第[一二三四五六七八九十百零〇\d]+\s*条)"
    )

    for document in documents:
        content = document["content"]
        parts = re.split(article_pattern, content)

        valid_parts = [
            part.strip()
            for part in parts
            if len(part.strip()) >= 20
        ]

        for index, part in enumerate(
            valid_parts,
            start=1,
        ):
            chunks.append(
                {
                    "chunk_id": (
                        f"{document['document_id']}"
                        f"-CHUNK-{index:03d}"
                    ),
                    "document_id": document["document_id"],
                    "title": document["title"],
                    "document_type": document[
                        "document_type"
                    ],
                    "source_url": document["source_url"],
                    "content": part,
                }
            )

    return chunks


def write_chunks_jsonl(
    chunks: list[dict[str, Any]],
    output_path: Path,
) -> None:
    """将文本片段逐行保存为JSONL文件。"""
    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with output_path.open(
        "w",
        encoding="utf-8",
    ) as file:
        for chunk in chunks:
            file.write(
                json.dumps(
                    chunk,
                    ensure_ascii=False,
                )
                + "\n"
            )