import re

from io import BytesIO

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from vehicle_valuation.model import (
    VehicleValuationReportDraft,
)

from pathlib import Path

from vehicle_valuation.docx_template import (
    fill_docx_template,
)


SECTION_FIELDS = [
    ("声明", "declaration"),
    ("摘要", "summary"),
    ("一、委托人及报告使用人", "client_and_users"),
    ("二、评估目的", "valuation_purpose"),
    ("三、评估对象和评估范围", "object_and_scope"),
    ("四、价值类型", "value_type"),
    ("五、评估基准日", "valuation_date"),
    ("六、评估依据", "valuation_basis"),
    ("七、评估方法", "valuation_method"),
    ("八、评估程序实施过程", "valuation_process"),
    ("九、评估假设", "assumptions"),
    ("十、评估结论", "conclusion"),
    ("十一、特别事项说明", "special_matters"),
    ("十二、报告使用限制", "usage_restrictions"),
]


def set_run_font(
    run,
    font_name: str,
    font_size: int,
    bold: bool = False,
) -> None:
    """设置Word文字的中英文字体、字号和粗体。"""

    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold

    run._element.get_or_add_rPr().rFonts.set(
        qn("w:eastAsia"),
        font_name,
    )


def configure_document(
    document: DocumentType,
) -> None:
    """设置A4纸张、页边距和正文样式。"""

    section = document.sections[0]

    section.page_width = Cm(21)
    section.page_height = Cm(29.7)

    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "STSong"
    normal_style.font.size = Pt(11)

    normal_style._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "STSong",
    )

    normal_format = normal_style.paragraph_format
    normal_format.space_after = Pt(6)
    normal_format.line_spacing = 1.25

    heading_style = document.styles["Heading 1"]
    heading_style.font.name = "Hiragino Sans GB"
    heading_style.font.size = Pt(16)
    heading_style.font.bold = True
    heading_style.font.color.rgb = RGBColor(
        46,
        116,
        181,
    )

    heading_style._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "Hiragino Sans GB",
    )

    heading_format = heading_style.paragraph_format
    heading_format.space_before = Pt(16)
    heading_format.space_after = Pt(8)
    heading_format.keep_with_next = True


def add_title(
    document: DocumentType,
    title: str,
) -> None:
    """添加报告标题和初稿提示。"""

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )
    title_paragraph.paragraph_format.space_after = Pt(12)

    title_run = title_paragraph.add_run(title)
    set_run_font(
        title_run,
        font_name="Hiragino Sans GB",
        font_size=22,
        bold=True,
    )

    draft_paragraph = document.add_paragraph()
    draft_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )
    draft_paragraph.paragraph_format.space_after = Pt(18)

    draft_run = draft_paragraph.add_run(
        "自动生成初稿——须经资产评估专业人员复核"
    )
    set_run_font(
        draft_run,
        font_name="STSong",
        font_size=10,
    )
    draft_run.font.color.rgb = RGBColor(
        156,
        101,
        0,
    )


def remove_inline_citations(
    text: str,
) -> str:
    """删除正文中的RAG片段编号，引用统一放在报告末尾。"""

    return re.sub(
        r"\s*\[DOC-\d+-CHUNK-\d+\]",
        "",
        text,
    )


def add_content(
    document: DocumentType,
    content: str,
) -> None:
    """将多段文字按换行拆成Word段落。"""

    cleaned_content = remove_inline_citations(
        content
    )

    for text in cleaned_content.splitlines():
        cleaned_text = text.strip()

        if not cleaned_text:
            continue

        paragraph = document.add_paragraph()
        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.LEFT
        )
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.25

        run = paragraph.add_run(cleaned_text)
        set_run_font(
            run,
            font_name="STSong",
            font_size=11,
        )


def add_missing_information(
    document: DocumentType,
    missing_information: list[str],
) -> None:
    """在报告末尾添加待补充资料清单。"""

    document.add_heading(
        "十三、待补充资料",
        level=1,
    )

    if not missing_information:
        add_content(
            document,
            "当前未记录待补充资料。",
        )
        return

    for number, item in enumerate(
        missing_information,
        start=1,
    ):
        paragraph = document.add_paragraph()

        run = paragraph.add_run(
            f"{number}. {item}"
        )
        set_run_font(
            run,
            font_name="STSong",
            font_size=11,
        )


def add_citations(
    document: DocumentType,
    cited_chunk_ids: list[str],
) -> None:
    """在报告末尾添加RAG证据片段编号。"""

    document.add_heading(
        "十四、RAG引用依据",
        level=1,
    )

    if not cited_chunk_ids:
        add_content(
            document,
            "本报告未记录RAG引用。",
        )
        return

    add_content(
        document,
        "本初稿引用的知识库片段编号如下：",
    )

    for chunk_id in cited_chunk_ids:
        paragraph = document.add_paragraph()

        run = paragraph.add_run(chunk_id)
        set_run_font(
            run,
            font_name="Consolas",
            font_size=9,
        )


def build_report_docx(
    draft: VehicleValuationReportDraft,
) -> bytes:
    """把结构化车辆评估初稿生成DOCX文件。"""

    document = Document()

    configure_document(document)
    add_title(document, draft.report_title)

    for heading, field_name in SECTION_FIELDS:
        document.add_heading(
            heading,
            level=1,
        )

        add_content(
            document,
            getattr(draft, field_name),
        )

    add_missing_information(
        document,
        draft.missing_information,
    )

    add_citations(
        document,
        draft.cited_chunk_ids,
    )

    output = BytesIO()
    document.save(output)

    return output.getvalue()


REPORT_TEMPLATE_PATH = Path(
    "templates/vehicle_valuation_report_template.docx"
)


def build_report_docx_from_template(
    draft: VehicleValuationReportDraft,
    fact_replacements: dict[str, str],
) -> bytes:
    """使用正式Word模板生成车辆评估报告初稿。"""

    replacements = dict(fact_replacements)

    replacements.update(
        {
            "{{DECLARATION}}": draft.declaration,
            "{{SUMMARY}}": draft.summary,
            "{{OBJECT_AND_SCOPE}}": (
                draft.object_and_scope
            ),
            "{{VALUE_TYPE}}": draft.value_type,
            "{{VALUATION_BASIS}}": (
                draft.valuation_basis
            ),
            "{{VALUATION_METHOD}}": (
                draft.valuation_method
            ),
            "{{VALUATION_PROCESS}}": (
                draft.valuation_process
            ),
            "{{ASSUMPTIONS}}": draft.assumptions,
            "{{CONCLUSION}}": draft.conclusion,
            "{{SPECIAL_MATTERS}}": (
                draft.special_matters
            ),
            "{{USAGE_RESTRICTIONS}}": (
                draft.usage_restrictions
            ),
        }
    )

    replacements = {
        placeholder: remove_inline_citations(value)
        for placeholder, value in replacements.items()
    }

    return fill_docx_template(
        REPORT_TEMPLATE_PATH,
        replacements,
    )